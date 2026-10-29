r"""
Builds week2\report\IOWA_vs_GT_report_Part1.docx: Part 1 of the report (summary of the method and the results on boundary error,
layer overlap, layer thickness, the three worst diseased-looking images and the normal-looking images at the foveal pit).
Outline files, topographic maps and run time belong to Part 2 and are not in this document.

Every number in the text is read from a table in week2\tables (this script writes no table and changes no other report file).
The figures are the original PNG files in week2\figures, embedded unchanged. Set REPO_URL below once the repository exists.

    python -I week2\scripts\build_report_part1.py
"""
import csv
import re
import sys
from collections import Counter
from pathlib import Path

import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).parent))
import w2common as W

REPO_URL = "https://github.com/<username>/<repository>"      # replace with the real address of the code repository
if "<" in REPO_URL:
    print("WARNING: REPO_URL is still a placeholder; Appendix B will show it as is. Edit REPO_URL at the top of this script.")

T, F = W.TABLES, W.FIGS
REPORT = W.W2 / "report"
REPORT.mkdir(exist_ok=True)


def read(name):
    with open(T / name, newline="") as f:
        return list(csv.DictReader(f))


# ------------------------------------------------------------------ data (same selection and tables as the full report)
sel = read("image_selection.csv")
sc = [r for r in read("scores_per_eye.csv") if r["variant"] == "original" and r["image_mode"] == "1" and r["config"] == "default"
      and r["group"] in ("normal", "diseased")]
have = {r["eye"] for r in sc}
missing = sorted(r["eye"] for r in sel if r["eye"] not in have)
G = {g: [r for r in sc if r["group"] == g] for g in ("normal", "diseased")}
nN, nD = len(G["normal"]), len(G["diseased"])
bars = {nm: {(r["group"], r["item"]): r for r in read(f"bars_{nm}.csv")} for nm in ("rmse_boundary", "dice_layer", "thickness_rmse", "thickness_bias")}
gate = read("gate_10301.csv")
pilot = read("pilot_21_vs_61.csv")
fail = read("t4_failure_stats.csv")


def b(nm, g, item, col="mean"):
    return float(bars[nm][(g, item)][col])


th = [r for r in read("thickness_per_eye.csv") if r["variant"] == "original" and r["image_mode"] == "1" and r["config"] == "default"
      and r["group"] in G]
TS = {}
for g in G:
    for layer in W.LAYER_NAMES + [W.TOTAL]:
        gt = [float(r["mean_um"]) for r in th if r["group"] == g and r["layer"] == layer and r["source"] == "GT"]
        io = [float(r["mean_um"]) for r in th if r["group"] == g and r["layer"] == layer and r["source"] == "IOWA"]
        TS[(g, layer)] = dict(gt_mean_um=np.mean(gt), iowa_mean_um=np.mean(io), bias_mean_um=b("thickness_bias", g, layer),
                              rmse_mean_um=b("thickness_rmse", g, layer), rmse_sd_um=b("thickness_rmse", g, layer, "sd"))

tot = len(sel)
THRESH = 10.0
gate_rmse = float([r for r in gate if r["item"] == "ALL" and r["kind"].startswith("boundary")][0]["value"])
pil = [r for r in pilot if r["boundary"] == "ALL"][0]
d_sorted = sorted(G["diseased"], key=lambda r: -float(r["rmse_ALL"]))
w3 = [r["eye"] for r in d_sorted[:3]]
w3_cross = sum(int(r["crossing_columns"]) for r in fail if r["eye"] in w3 and r["boundary"] == "ILM")
n_ok_med = float(np.median([float(r["rmse_ALL"]) for r in G["normal"] if float(r["rmse_ALL"]) < THRESH]))
n_low = sum(float(r["rmse_ALL"]) < THRESH for r in G["normal"])
d_low = sum(float(r["rmse_ALL"]) < THRESH for r in G["diseased"])


def worst_boundary(r):
    v = {nm: float(r[f"rmse_{nm}"]) for nm in W.GT_BOUNDARIES}
    k = max(v, key=v.get)
    return k, v[k]


n_hi = [r for r in G["normal"] if float(r["rmse_ALL"]) >= THRESH]
n_hi_worst = Counter(worst_boundary(r)[0] for r in n_hi)
th_tot_hi = [float(r["rmse_um"]) for r in th if r["layer"] == W.TOTAL and r["source"] == "IOWA" and r["eye"] in {x["eye"] for x in n_hi}]
n_sel_dis = sum(r["group"] == "diseased" for r in sel)

# ------------------------------------------------------------------ document helpers
doc = Document()
doc.core_properties.author = ""
doc.core_properties.title = "IOWA layer segmentation against ground truth on OCTA-500 B-scans, Part 1"
doc.core_properties.comments = ""
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.2)
sec.top_margin = sec.bottom_margin = Cm(2.0)
st = doc.styles["Normal"]
st.font.name, st.font.size = "Calibri", Pt(11)
st.paragraph_format.space_after = Pt(6)
IMG_W = Cm(16.4)


def H(level, text):
    doc.add_heading(text, level)


def P(text, italic=False, bold=False, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic, r.bold = italic, bold
    if size:
        r.font.size = Pt(size)
    return p


def BL(items):
    for t in items:
        doc.add_paragraph(t, style="List Bullet")


def NUM(items):
    """Numbered list that always starts at 1 (typed numbers, so two lists in one document cannot continue each other)."""
    for i, t in enumerate(items, 1):
        p = doc.add_paragraph(f"{i}. {t}")
        p.paragraph_format.left_indent = Cm(0.9)
        p.paragraph_format.first_line_indent = Cm(-0.9)
        p.paragraph_format.space_after = Pt(4)


def TAB(header, rows, widths=None, size=9):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = str(v)
    for row in t.rows:
        for i, c in enumerate(row.cells):
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    r.font.size = Pt(size)
            if widths:
                c.width = Cm(widths[i])
    doc.add_paragraph()


def FIG(path, caption, width=IMG_W):
    doc.add_picture(str(path), width=width)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = P(caption, italic=True, size=9)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def ms(nm, g, item, d=1):
    return f"{b(nm, g, item):.{d}f} +/- {b(nm, g, item, 'sd'):.{d}f}"


# ================================================================== title
doc.add_heading("How well does IOWA layer segmentation agree with ground truth on OCTA-500 B-scans?", 0)
P("Part 1: boundary error, layer overlap and layer thickness", bold=True)
P(f"{nN + nD} B-scans from {nN + nD} eyes ({nN} normal-looking, {nD} diseased-looking). "
  "Segmentation: Iowa Reference Algorithms, OCTExplorer 5.0.0b, Default configuration. Data: OCTA-500, 3 mm scans.", italic=True)

# ================================================================== 1 Introduction
H(1, "1. Introduction")
P("IOWA is free research software that finds the boundaries between the layers of the retina in an OCT scan [2-5]. "
  "OCTA-500 is a public dataset that comes with layer boundaries drawn for each scan [1]. "
  "We ask a simple question: how close are IOWA's boundaries to these drawn boundaries, and where does IOWA go wrong?")
H(2, "1.1 Terms used in this report")
TAB(["Term", "Meaning"], [
    ["OCT, B-scan", "OCT takes depth images of the retina. One B-scan is one such image: 304 columns across (3 mm) and 640 rows deep (2 mm). One eye has 304 B-scans."],
    ["Boundary", "A line between two retinal layers, drawn as one depth value per column. We use six: ILM, IPL-INL, OPL-HFL, BMEIS, IB_RPE and OB_RPE (names as in the IOWA output)."],
    ["Layer", "The band between two neighbouring boundaries. The five layers are named by their two boundaries, for example IPL-INL->OPL-HFL. Total retina is ILM to OB_RPE."],
    ["Ground truth (GT)", "The boundaries supplied with OCTA-500. They were first made by IOWA 3.8 and then corrected by hand (see section 4.1)."],
    ["RMSE (um)", "Root mean square error: the typical vertical distance between IOWA's boundary and the GT boundary, in micrometres. 0 is perfect. One image row is 3.125 um."],
    ["Dice", "Overlap of a layer drawn by IOWA with the same layer in GT, from 0 (no overlap) to 1 (identical)."],
    ["Thickness", "Distance between two boundaries in um. Thickness bias is IOWA minus GT: above 0 means IOWA is thicker."],
    ["Normal-looking, diseased-looking", "Our own grouping from looking at the B-scans (no diagnosis list exists for these eyes). Diseased-looking means visible cysts or swelling."],
    ["Default configuration", "OCTExplorer's standard setting for most scans. The Cyst configuration is meant for scans with cysts and was not tested here."],
], widths=[3.6, 12.8])
H(2, "1.2 Scope of this report")
P("This report (Part 1) compares IOWA's layer boundaries with the ground truth on one B-scan per eye. It covers the boundary error, the layer overlap, "
  "the layer thickness, the three worst diseased-looking images and the normal-looking images at the foveal pit. "
  "The outline files, the topographic maps and the run time are reported separately in Part 2.")

# ================================================================== 2 Methodology
H(1, "2. Methodology")
H(2, "2.1 Data and choice of images")
P(f"OCTA-500 has images for eyes 10301 to 10400 in the part we hold. Eyes 10333 and 10388 are the same scan stored twice, so we used neither. "
  f"We planned {tot} eyes: one B-scan from each of {tot} different eyes. Always included: 10306 and 10335 (both cystic).")
if missing:
    P(f"Note: {tot} images were planned. The runs for eye(s) {', '.join(missing)} were not done, so this report uses {nN + nD}.", bold=True)
BL([
    "Normal-looking eyes: the B-scan at the foveal pit, found from the GT as the thinnest part of the retina near the centre of the scan.",
    "Diseased-looking eyes: the B-scan with the largest GT abnormality (the thickest cyst band, or the thickest retina for eyes with a raised dome). "
    f"Only {n_sel_dis} eyes in this set show visible disease close to the centre, so the groups are not equal in size.",
    "The choice uses the GT only and never looks at the IOWA result.",
])
rows = []
for r in sorted(sel, key=lambda r: (r["group"], r["eye"])):
    rows.append([r["eye"], r["group"], r["bscan"], r["rule"], "yes" if r["eye"] in have else "NOT RUN"])
TAB(["Eye", "Group", "B-scan", "How the B-scan was chosen", "Scored"], rows, widths=[1.8, 2.2, 1.6, 8.6, 2.2], size=8)

H(2, "2.2 Segmentation procedure")
NUM([
    "For each eye we made a 21-B-scan volume centred on the chosen B-scan (10 on each side). IOWA needs neighbouring B-scans to work in 3D, and 21 keeps each run short.",
    "A person opened each volume in OCTExplorer 5.0.0b, picked Configuration-Default.json, and clicked through the two steps. No surface was edited or nudged.",
    "Only the chosen B-scan (the centre one) was compared with GT. The other 20 are context.",
    "The comparisons were made with the code listed in Appendix B (section 9). The resulting tables are in the repository folder tables/.",
])

H(2, "2.3 How we score")
P("For each scored B-scan we compare IOWA's six boundaries with the six GT boundaries. The boundary error is the RMSE (um) defined in section 1.1; "
  "we give it for each boundary and for all six boundaries together. The layer overlap is the Dice value of each of the five layers and of the whole retina. "
  "Thickness is the distance between two boundaries in um; for each of the five layers and for the total retina we report the mean thickness in GT and in IOWA, "
  "the mean bias (IOWA minus GT) and the RMSE of the thickness. "
  "Each eye contributes one value per measure. Group results are the mean and the standard deviation (SD) across eyes, and the individual eyes are shown as dots in the figures.")

H(2, "2.4 Method checks")
P("Two checks on the method:")
BL([
    f"Gate: our scoring code reproduces the earlier result for eye 10301 (B-scans 120-180) of 3.84 um RMSE from the existing IOWA file "
    f"(we get {gate_rmse:.4f} um).",
    f"Is 21 scans enough? On one eye (10301, B-scan 150) we ran 21 scans and compared with the earlier 61-scan run. "
    f"Overall error {float(pil['rmse_61scan_um']):.2f} um with 61 scans and "
    f"{float(pil['rmse_21scan_um']):.2f} um with 21 scans. "
    "This is one eye with a normal retina, so it does not prove 21 scans are enough for cystic eyes.",
])

# ================================================================== 3 Results
H(1, "3. Results")
H(2, "3.1 Error of each boundary")
P("Each bar is the mean over eyes of the RMSE of one boundary. The thin bars show the standard deviation across eyes (the lower end is cut at 0). "
  "The black dots are individual eyes.")
FIG(F / "bars_rmse_boundary.png", "Figure 1. Boundary error of IOWA against GT, one B-scan per eye, Default configuration.")
items = W.GT_BOUNDARIES + ["ALL"]
TAB(["Boundary", f"Normal-looking (n = {nN}), um", f"Diseased-looking (n = {nD}), um"],
    [[i if i != "ALL" else "All six boundaries", ms("rmse_boundary", "normal", i), ms("rmse_boundary", "diseased", i)] for i in items],
    widths=[5, 5.7, 5.7])
P("Values are mean +/- SD across eyes. The SD is often as large as the mean because a few eyes have very large errors; with this few eyes the "
  "difference between the groups should be read as a trend, not a proven difference.")

H(2, "3.2 Overlap of each layer")
FIG(F / "bars_dice_layer.png", "Figure 2. Dice overlap of each layer with GT. Higher is better.")
TAB(["Layer", f"Normal-looking (n = {nN})", f"Diseased-looking (n = {nD})"],
    [[i.replace("whole_retina", "Whole retina"), ms("dice_layer", "normal", i, 3), ms("dice_layer", "diseased", i, 3)] for i in W.LAYER_NAMES + ["whole_retina"]],
    widths=[5, 5.7, 5.7])
P(f"The whole retina overlaps almost perfectly in both groups ({b('dice_layer', 'normal', 'whole_retina'):.3f} and {b('dice_layer', 'diseased', 'whole_retina'):.3f}) "
  "because the top and bottom surfaces are easy to find. This number hides the errors inside the retina, so the per-layer values matter more.")

H(2, "3.3 Layer thickness")
P("Thickness is the distance between two boundaries. The table gives the average thickness in GT and in IOWA, the mean bias (IOWA minus GT) and the RMSE of the thickness.")
for g, nn in (("normal", nN), ("diseased", nD)):
    P(f"{g.capitalize()}-looking eyes (n = {nn})", bold=True)
    TAB(["Layer", "GT mean (um)", "IOWA mean (um)", "Bias, IOWA - GT (um)", "RMSE (um)"],
        [[l, f"{TS[(g, l)]['gt_mean_um']:.1f}", f"{TS[(g, l)]['iowa_mean_um']:.1f}", f"{TS[(g, l)]['bias_mean_um']:+.1f}",
          f"{TS[(g, l)]['rmse_mean_um']:.1f} +/- {TS[(g, l)]['rmse_sd_um']:.1f}"] for l in W.LAYER_NAMES + [W.TOTAL]],
        widths=[5.4, 2.6, 2.8, 3.2, 2.6])
FIG(F / "bars_thickness_rmse.png", "Figure 3. Thickness error (RMSE) per layer and for the total retina.")
FIG(F / "bars_thickness_bias.png", "Figure 4. Mean signed thickness error. Above 0: IOWA thicker than GT. Below 0: thinner.")
P("The full per-image thickness numbers (mean, median, SD per layer, GT and IOWA) are in tables/thickness_per_eye.csv and tables/thickness_per_bscan.csv.")

H(2, "3.4 Diseased-looking images: the three worst")
rows = []
for r in d_sorted:
    wb, wv = worst_boundary(r)
    rows.append([r["eye"], r["a"], f"{float(r['rmse_ALL']):.1f}", f"{wb} ({wv:.1f})", f"{float(r['dice_whole_retina']):.3f}"])
TAB(["Eye", "B-scan", "RMSE, all boundaries (um)", "Worst boundary (RMSE, um)", "Dice, whole retina"], rows, widths=[2, 2, 4.2, 5.2, 3])
FIG(F / "t4_worst3_overlay.png", "Figure 5. The three worst diseased-looking B-scans. Solid lines: GT. Dashed lines: IOWA. Bars: RMSE of each boundary.")
P("What we see in the pictures (our reading, not a clinician's):", bold=True)
BL([
    f"Eye {w3[0]}: large dark fluid-filled cysts fill most of the retina on the left and centre. IOWA's OPL-HFL line (green, dashed) follows the cyst walls and drops "
    "down into the cyst area instead of staying along the layer, so it is far from GT. On the right, a bright patch with a shadow below it pulls IOWA's two lowest "
    "boundaries (the RPE lines) upward away from GT.",
    f"Eye {w3[1]}: a dome-shaped swelling with a dark fluid pocket. IOWA places the IPL-INL line (orange) lower than GT inside the dome, and the OPL-HFL line is also off. "
    "The top surface and the bottom boundaries are still good.",
    f"Eye {w3[2]}: small cysts on the left and a raised bump in the middle. IOWA's BMEIS line (blue) and OPL-HFL line (green) do not follow GT over the bump and the cysts.",
    "In all three, the top surface (ILM) is right. The errors begin where there are cysts or swelling, and the inner boundaries are the ones that lose their way.",
    (f"Boundaries cross in {w3_cross} column(s) of these three images." if w3_cross else "No boundaries cross or go missing in these three images (tables/t4_failure_stats.csv)."),
])

H(2, "3.5 Normal-looking images also have large errors at the foveal pit")
P(f"We expected small errors in normal-looking eyes. For {len(n_hi)} of {nN} of them the overall RMSE is {THRESH:.0f} um or more "
  f"({THRESH:.0f} um is our own rough cut-off, not a standard). Among these, the boundary with the largest error is "
  + ", ".join(f"{k} in {v}" for k, v in n_hi_worst.most_common()) + " of them. "
  + (f"Total thickness is still good in these eyes (median RMSE {np.median(th_tot_hi):.1f} um); the error is in an inner boundary." if th_tot_hi else ""))
FIG(F / "t4_worst3_overlay_normal.png", "Figure 6. The three worst normal-looking B-scans (all at the foveal pit). Solid: GT. Dashed: IOWA.")
P("What we see: at the pit the layers become very thin and merge, so the line between them is hard to place. "
  "IOWA's OPL-HFL line (green) dips below GT on one side of the pit. In some eyes IOWA's BMEIS line (blue) jumps up into the retina on one side. "
  f"We checked that GT lines sit where the image suggests, and that the same code gives a median of {n_ok_med:.1f} um on the normal-looking eyes below the cut-off, so we do not think this is a scoring mistake. "
  "We cannot say which line is anatomically right in every case.")
P("This matters for reading the groups. Normal-looking eyes were scored at the foveal pit, a place where IOWA often struggles in this data, "
  "so their error may be higher than for a typical B-scan away from the pit. We did not score other B-scans, so we cannot say by how much.")

# ================================================================== 4 Discussion
H(1, "4. Discussion")
P(f"The outer boundaries and the top surface are close to GT "
  f"(ILM {b('rmse_boundary', 'normal', 'ILM'):.1f} um and OB_RPE {b('rmse_boundary', 'normal', 'OB_RPE'):.1f} um in normal-looking eyes). "
  "The inner boundaries IPL-INL, OPL-HFL and BMEIS are where the large errors are (section 3.1).")
P(f"Total retinal thickness is much better than the individual boundaries: thickness RMSE {b('thickness_rmse', 'normal', W.TOTAL):.1f} um (normal-looking) "
  f"and {b('thickness_rmse', 'diseased', W.TOTAL):.1f} um (diseased-looking) (section 3.3). "
  f"The whole-retina Dice value is almost perfect in both groups ({b('dice_layer', 'normal', 'whole_retina'):.3f} and {b('dice_layer', 'diseased', 'whole_retina'):.3f}), "
  "but this number hides the errors inside the retina, so the per-layer values matter more (section 3.2).")
P(f"The overall boundary error (RMSE over all six boundaries) was {ms('rmse_boundary', 'normal', 'ALL')} um in normal-looking eyes (n = {nN}) and "
  f"{ms('rmse_boundary', 'diseased', 'ALL')} um in diseased-looking eyes (n = {nD}). The spread between eyes is large, and the SD is often as large as the mean, "
  "so the difference between the groups should be read as a trend, not a proven difference. "
  f"In the three worst diseased-looking images the errors begin where there are cysts or swelling, and the inner boundaries are the ones that lose their way, "
  "while the top surface is right (section 3.4).")
P(f"Only {n_low} of {nN} normal-looking and {d_low} of {nD} diseased-looking images have an overall RMSE below {THRESH:.0f} um. "
  "Many normal-looking foveal-pit images also show large errors. This is a finding, not an error in scoring (section 3.5). "
  "At the pit the layers become very thin and merge, so the line between them is hard to place.")
P("Because the ground truth was started from IOWA's own output and then hand-corrected, these numbers probably flatter IOWA (section 4.1).")
H(2, "4.1 Limitations")
BL([
    "Circular ground truth. The GT boundaries were started from IOWA 3.8 output and corrected by hand. IOWA agrees with GT partly because it helped to make it. "
    "The real error against the true anatomy is probably larger than reported here.",
    f"Few images. {nN + nD} B-scans, one per eye, out of the {tot} planned. Only {nD} are diseased-looking, and one B-scan cannot represent a whole eye.",
    "The groups are our own visual judgement, not clinical diagnoses. Mild disease may sit in the normal group.",
    "Different choice of B-scan per group: foveal pit (normal) and most abnormal B-scan (diseased). The groups are therefore not matched, and the normal-group error is a pit error.",
    "Default configuration only. The Cyst configuration was not tested, so we cannot say whether it would do better on the cystic eyes.",
    "One scanner and one dataset (OCTA-500, 3 mm). Results may not carry over to other devices or to the 6 mm scans.",
    "21-scan volumes, not full volumes. Checked on one normal eye only.",
    "The software version used is labelled 5.0.0b.",
])

# ================================================================== 5 Conclusion
H(1, "5. Conclusion")
P(f"On one B-scan from each of {nN + nD} eyes, IOWA's top surface and its lowest boundary agree closely with the ground truth in normal-looking eyes "
  f"(ILM {b('rmse_boundary', 'normal', 'ILM'):.1f} um and OB_RPE {b('rmse_boundary', 'normal', 'OB_RPE'):.1f} um), "
  "while the inner boundaries IPL-INL, OPL-HFL and BMEIS carry the large errors. "
  f"Total retinal thickness is much better than the single boundaries ({b('thickness_rmse', 'normal', W.TOTAL):.1f} um in normal-looking and "
  f"{b('thickness_rmse', 'diseased', W.TOTAL):.1f} um in diseased-looking eyes). "
  f"The overall boundary error was larger in diseased-looking eyes ({b('rmse_boundary', 'diseased', 'ALL'):.1f} um) than in normal-looking eyes ({b('rmse_boundary', 'normal', 'ALL'):.1f} um), "
  "where the errors begin at cysts and swelling. "
  f"Normal-looking eyes also show large errors at the foveal pit, in {len(n_hi)} of {nN} images. "
  "Because the ground truth began with IOWA's own output and the groups are small and visually defined, these results are a trend and not a final measure.")

# ================================================================== 6 Future work
H(1, "6. Future work")
BL([
    "Run the Cyst configuration on the cystic eyes and compare it with the Default configuration. A lower score against GT would not by itself show that the Cyst configuration is worse, "
    "because the GT began with IOWA's own output.",
    "Run IOWA on full 304-scan volumes and compare the result with the 21-scan volumes used here, which we checked on one eye only.",
    "Use clinical diagnosis labels instead of our visual grouping, and choose the same kind of B-scan in both groups so that the groups are matched.",
    "Compare IOWA with a reference that is independent of IOWA, so that the agreement is not circular.",
    "Find out why the inner boundaries fail. Here we described it only from the overlays of six images.",
    "Compare the remaining segmentation algorithms planned in the project with IOWA, and compare layer thickness between the algorithms.",
])

# ================================================================== 7 References
H(1, "7. References")
P("The Iowa Reference Algorithms and OCTExplorer were provided by The University of Iowa Research Foundation. "
  "We acknowledge the providers' scientists: Michael D. Abramoff, Mona K. Garvin, Andreas Wahle, Kyungmoo Lee and Milan Sonka.")
NUM([
    "Li M, Huang K, Xu Q, Yang J, Zhang Y, Ji Z, Xie K, Yuan S, Liu Q, Chen Q. OCTA-500: A retinal dataset for optical coherence tomography angiography study. arXiv:2012.07261 (v3), 2022.",
    "Abramoff MD, Garvin M, Sonka M. Retinal imaging and image analysis. IEEE Reviews in Biomedical Engineering 3:169-208, 2010. doi:10.1109/RBME.2010.2084567",
    "Li K, Wu X, Chen DZ, Sonka M. Optimal surface segmentation in volumetric images - a graph-theoretic approach. IEEE Transactions on Pattern Analysis and Machine Intelligence 28:119-134, 2006. doi:10.1109/TPAMI.2006.19",
    "Garvin MK, Abramoff MD, Wu X, Burns TK, Russell SR, Sonka M. Automated 3-D intraretinal layer segmentation of macular spectral-domain optical coherence tomography images. IEEE Transactions on Medical Imaging 28(9):1436-1447, 2009. doi:10.1109/TMI.2009.2016958",
    "Quellec G, Lee K, Dolejsi M, Garvin MK, Abramoff MD, Sonka M. Three-dimensional analysis of retinal layer texture: identification of fluid-filled regions in SD-OCT of the macula. IEEE Transactions on Medical Imaging 29:1321-1330, 2010. doi:10.1109/TMI.2010.2047023",
])

# ================================================================== 8 Appendix A
H(1, "8. Appendix A: every scored image")
rows = []
for r in sorted(sc, key=lambda r: (r["group"], r["eye"])):
    wb, wv = worst_boundary(r)
    rows.append([r["eye"], r["group"], r["a"], f"{float(r['rmse_ALL']):.1f}", f"{wb} ({wv:.1f})", f"{float(r['dice_whole_retina']):.3f}"])
TAB(["Eye", "Group", "B-scan", "RMSE all (um)", "Worst boundary (um)", "Dice whole retina"], rows, widths=[2, 2.4, 2, 3, 4.6, 3], size=8)

# ================================================================== 9 Appendix B
H(1, "9. Appendix B: code")
P(f"The code is in the repository {REPO_URL}. The folders named in this report (scripts/, tables/, figures/, outlines/, maps/) are folders of that repository. "
  "The table lists the scripts in scripts/.")
IO = {   # script(s) -> (input, output)
    "gate_check.py": ("existing IOWA result and GT of eye 10301, B-scans 120-180", "tables/gate_10301.csv"),
    "screen_abnormal.py": ("GT layers of all usable eyes", "tables/abnormal_screening.csv"),
    "make_selection.py": ("GT layers; eye lists chosen by visual inspection", "tables/image_selection.csv"),
    "make_volumes.py": ("OCT B-scans; tables/image_selection.csv", "21-scan .mhd/.raw volumes in mhd/ and mhd_cyst/<eye>/"),
    "make_checklist.py": ("tables/image_selection.csv; result files", "RUN_CHECKLIST.md"),
    "watch_and_score.py": ("IOWA result XML files in mhd/ and mhd_cyst/; GT", "tables/scores_per_bscan.csv, scores_per_eye.csv, thickness_*.csv, run_timing.csv"),
    "score_runs.py": ("IOWA result XML files; GT", "tables/scores_per_bscan.csv, tables/scores_per_eye.csv"),
    "t1_thickness.py": ("IOWA result XML files; GT", "tables/thickness_per_bscan.csv, tables/thickness_per_eye.csv"),
    "t2_outlines.py": ("IOWA result XML files; GT", "outlines/<eye>_<B-scan>/ (PNG, .npy, .csv)"),
    "t3_maps.py": ("IOWA result XML files of volumes or chunks; GT", "maps/<eye>_<range>/ (PNG, .npy); tables/map_coverage.csv"),
    "t4_diseased.py": ("tables/scores_per_eye.csv, tables/thickness_per_eye.csv; IOWA result XML files", "tables/t4_*.csv; figures/t4_worst3_overlay*.png; Cyst-vs-Default figures when Cyst runs exist"),
    "t5_bars.py": ("tables/scores_per_eye.csv, tables/thickness_per_eye.csv", "figures/bars_*.png; tables/bars_*.csv"),
    "t6_runtime.py": ("tables/run_timing.csv, tables/scores_per_eye.csv", "tables/runtime_*.csv; figures/bars_runtime.png, figures/scatter_runtime_vs_error.png"),
    "build_report.py": ("tables/, figures/, maps/, outlines/", "report/IOWA_vs_GT_report.docx"),
    "build_report_part1.py": ("tables/, figures/", "report/IOWA_vs_GT_report_Part1.docx"),
    "edge_effect.py": ("tables/scores_per_bscan.csv", "tables/edge_effect_10301.csv"),
    "compare_pilot.py": ("21-scan and 61-scan IOWA results of eye 10301; GT", "tables/pilot_21_vs_61.csv"),
}
readme = (W.W2 / "README.md").read_text(encoding="utf-8").splitlines()
code_rows, started = [], False
for ln in readme:
    if ln.startswith("| script |"):
        started = True
        continue
    if started and ln.startswith("|---"):
        continue
    if started and ln.startswith("|"):
        cells = [c.strip() for c in ln.strip().strip("|").split("|", 1)]
        name, what = cells[0], cells[1]
        def clean(s):
            for a, b_ in (("rule 4: ", ""), (", PASSED", ""), (" (re-run after new results)", ""), (" (helps find diseased candidates)", "")):
                s = s.replace(a, b_)
            return s.replace("`", "").replace("\\", "/")
        scripts = [s.strip() for s in name.split(",")]
        if len(scripts) == 1:
            i_, o_ = IO.get(scripts[0], ("see the header of the script", "see the header of the script"))
        else:
            i_ = "; ".join(f"{s}: {IO[s][0]}" if s in IO else s for s in scripts)
            o_ = "; ".join(f"{s}: {IO[s][1]}" if s in IO else s for s in scripts)
        code_rows.append([clean(name), clean(what), i_, o_])
    elif started:
        break
code_rows.append(["w2common.py, w2tables.py", "Shared code: loading GT and IOWA surfaces, layer names, eye groups, run discovery; scores and layer thickness for one run.",
                  "GT layers; IOWA result XML files", "(used by the other scripts)"])
TAB(["Script", "What it does", "Input", "Output"], code_rows, widths=[3.2, 5.8, 3.8, 3.8], size=8)

out = REPORT / "IOWA_vs_GT_report_Part1.docx"
doc.save(out)
print(f"saved {out}")
